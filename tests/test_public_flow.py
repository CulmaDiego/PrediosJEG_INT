from __future__ import annotations

import io

from openpyxl import load_workbook
from sqlalchemy import select

from app.models import DrawingVersion, Person, PublicLink, TileWindow
from app.security import public_csrf, sha256, stable_public_token
from tests.conftest import login


def square(lon: float, lat: float, delta: float = 0.0002) -> dict:
    return {"type": "Polygon", "coordinates": [[[lon - delta, lat - delta], [lon + delta, lat - delta], [lon + delta, lat + delta], [lon - delta, lat + delta], [lon - delta, lat - delta]]]}


def test_valid_polygon_creates_immutable_current_version(test_app):
    app, client, token, person_id = test_app
    csrf = public_csrf(token, app.state.settings)
    first = client.post(f"/p/{token}/submit", json={"csrf": csrf, "geometry": square(-75.031, 6.294)})
    second = client.post(f"/p/{token}/submit", json={"csrf": csrf, "geometry": square(-75.031, 6.294, 0.00015)})
    assert first.status_code == 200
    assert second.status_code == 200
    assert first.json()["redirect_url"].endswith("/recibido")
    assert second.json()["redirect_url"].endswith("/recibido")
    success = client.get(second.json()["redirect_url"])
    assert success.status_code == 200
    assert "Respuesta recibida" in success.text
    with app.state.Session() as db:
        drawings = db.scalars(select(DrawingVersion).where(DrawingVersion.person_id == person_id).order_by(DrawingVersion.version)).all()
        assert [item.version for item in drawings] == [1, 2]
        assert [item.is_current for item in drawings] == [False, True]


def test_stable_link_reset_clears_versions_without_changing_url(test_app):
    app, client, _, person_id = test_app
    token = stable_public_token(person_id, app.state.settings)
    csrf = public_csrf(token, app.state.settings)
    assert client.post(f"/p/{token}/submit", json={"csrf": csrf, "geometry": square(-75.031, 6.294)}).status_code == 200
    admin_csrf = login(client)
    detail_before = client.get(f"/admin/personas/{person_id}")
    public_url = f"http://testserver/p/{token}"
    assert public_url in detail_before.text
    assert "Copiar enlace" in detail_before.text
    assert "Reiniciar respuesta" in detail_before.text
    reset = client.post(f"/admin/personas/{person_id}/reiniciar", data={"csrf_token": admin_csrf}, follow_redirects=False)
    assert reset.status_code == 303
    detail_after = client.get(f"/admin/personas/{person_id}")
    assert public_url in detail_after.text
    with app.state.Session() as db:
        person = db.get(Person, person_id)
        drawings = db.scalars(select(DrawingVersion).where(DrawingVersion.person_id == person_id)).all()
        assert drawings == []
        assert person.status == "PENDIENTE"
        assert person.drawing_version == 0
        assert person.responded_at is None


def test_bulk_links_and_exports_use_stable_links_and_latest_version(test_app):
    app, client, token, person_id = test_app
    csrf = public_csrf(token, app.state.settings)
    client.post(f"/p/{token}/submit", json={"csrf": csrf, "geometry": square(-75.031, 6.294)})
    client.post(f"/p/{token}/submit", json={"csrf": csrf, "geometry": square(-75.031, 6.294, 0.00015)})
    login(client)

    links = client.get("/admin/enlaces")
    assert links.status_code == 200
    links_sheet = load_workbook(io.BytesIO(links.content), read_only=True).active
    link_rows = list(links_sheet.iter_rows(values_only=True))
    assert link_rows[0] == ("NOMBRE", "CELULAR", "MUNICIPIO", "ENLACE")
    assert link_rows[1][3] == f"http://testserver/p/{stable_public_token(person_id, app.state.settings)}"

    report = client.get("/admin/reportes/xlsx")
    assert report.status_code == 200
    report_sheet = load_workbook(io.BytesIO(report.content), read_only=True).active
    report_rows = list(report_sheet.iter_rows(values_only=True))
    assert len(report_rows) == 2
    assert report_rows[1][4] == 2


def test_self_intersection_and_far_polygon_are_rejected(test_app):
    app, client, token, _ = test_app
    csrf = public_csrf(token, app.state.settings)
    bowtie = {"type": "Polygon", "coordinates": [[[-75.0312, 6.2938], [-75.0308, 6.2942], [-75.0312, 6.2942], [-75.0308, 6.2938], [-75.0312, 6.2938]]]}
    assert client.post(f"/p/{token}/submit", json={"csrf": csrf, "geometry": bowtie}).status_code == 422
    assert client.post(f"/p/{token}/submit", json={"csrf": csrf, "geometry": square(-75.00, 6.32)}).status_code == 422


def test_tile_from_another_person_cannot_be_read(test_app, tmp_path):
    app, client, token_one, person_one = test_app
    with app.state.Session() as db:
        municipality_id = db.get(Person, person_one).municipality_id
        person_two = Person(municipality_id=municipality_id, name="Otra", phone="3009999999", reference_lon=-75.031, reference_lat=6.294)
        db.add(person_two)
        db.flush()
        token_two = "b" * 43
        db.add(PublicLink(person_id=person_two.id, token_hash=sha256(token_two)))
        tile_dir = app.state.settings.data_dir / "tiles" / "first" / "18" / "76440"
        tile_dir.mkdir(parents=True)
        (tile_dir / "130000.png").write_bytes(b"not-a-real-png")
        db.add(TileWindow(person_id=person_one, relative_dir="first", min_zoom=18, max_zoom=18, status="LISTO"))
        db.commit()
    assert client.get(f"/p/{token_one}/tiles/18/76440/130000.png").status_code == 200
    assert client.get(f"/p/{token_two}/tiles/18/76440/130000.png").status_code == 404
