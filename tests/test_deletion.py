from __future__ import annotations

import io
import json

from openpyxl import load_workbook
from sqlalchemy import func, select

from app.models import AuditLog, DrawingVersion, ImportBatch, Job, Municipality, Person, PublicLink, TileWindow
from app.security import public_csrf, sha256
from tests.conftest import login
from tests.test_public_flow import square


def test_individual_delete_removes_database_links_files_and_exports(test_app):
    app, client, token, person_id = test_app
    settings = app.state.settings
    assert client.post(
        f"/p/{token}/submit",
        json={"csrf": public_csrf(token, settings), "geometry": square(-75.031, 6.294)},
    ).status_code == 200
    tile = settings.data_dir / "tiles" / "delete-me" / "18" / "1" / "2.png"
    tile.parent.mkdir(parents=True)
    tile.write_bytes(b"tile")
    imported = settings.data_dir / "imports" / "delete-me.xlsx"
    imported.write_bytes(b"source")
    old_shape = settings.data_dir / "exports" / "old" / "predios_shapefile.zip"
    old_shape.parent.mkdir(parents=True)
    old_shape.write_bytes(b"old-export")
    with app.state.Session() as db:
        person = db.get(Person, person_id)
        db.add(TileWindow(person_id=person_id, relative_dir="delete-me", min_zoom=18, max_zoom=18, status="LISTO"))
        db.add(Job(job_type="PREPARE_WINDOW", payload_json=json.dumps({"person_id": person_id}), status="COMPLETADO"))
        db.add(
            ImportBatch(
                municipality_id=person.municipality_id,
                stored_name=imported.name,
                original_name="personas.xlsx",
                status="CONFIRMADO",
                rows_json=json.dumps([{"phone": person.phone, "name": person.name}]),
                errors_json="[]",
                row_count=1,
            )
        )
        db.commit()

    csrf = login(client)
    response = client.post(
        f"/admin/personas/{person_id}/eliminar",
        data={"csrf_token": csrf},
        follow_redirects=False,
    )
    assert response.status_code == 303
    assert response.headers["location"] == "/admin/personas?eliminadas=1"
    assert client.get(f"/p/{token}").status_code == 404
    assert client.get(f"/admin/personas/{person_id}").status_code == 404
    with app.state.Session() as db:
        assert db.get(Person, person_id) is None
        assert db.scalar(select(func.count()).select_from(PublicLink).where(PublicLink.person_id == person_id)) == 0
        assert db.scalar(select(func.count()).select_from(DrawingVersion).where(DrawingVersion.person_id == person_id)) == 0
        assert db.get(TileWindow, person_id) is None
        assert db.scalar(select(func.count()).select_from(AuditLog).where(AuditLog.person_id == person_id)) == 0
        assert db.scalar(select(func.count()).select_from(Job).where(Job.payload_json.contains(str(person_id)))) == 0
        assert db.scalar(select(func.count()).select_from(ImportBatch).where(ImportBatch.stored_name == imported.name)) == 0
    assert not tile.parents[2].exists()
    assert not imported.exists()
    assert not old_shape.exists()

    links = load_workbook(io.BytesIO(client.get("/admin/enlaces").content), read_only=True).active
    report = load_workbook(io.BytesIO(client.get("/admin/reportes/xlsx").content), read_only=True).active
    assert list(links.iter_rows(values_only=True)) == [("NOMBRE", "CELULAR", "MUNICIPIO", "ENLACE")]
    assert list(report.iter_rows(values_only=True)) == [("PERSONA_ID", "NOMBRE", "CELULAR", "MUNICIPIO", "VERSION", "AREA_M2", "GEOJSON")]


def test_selected_and_municipality_delete_keep_unselected_and_municipality(test_app):
    app, client, token, original_id = test_app
    with app.state.Session() as db:
        municipality = db.scalar(select(Municipality).where(Municipality.name == "San Rafael"))
        selected = Person(municipality_id=municipality.id, name="Seleccionada", phone="3000000002")
        remaining = Person(municipality_id=municipality.id, name="Restante", phone="3000000003")
        db.add_all([selected, remaining])
        db.flush()
        db.add_all(
            [
                PublicLink(person_id=selected.id, token_hash=sha256("selected-token")),
                PublicLink(person_id=remaining.id, token_hash=sha256("remaining-token")),
            ]
        )
        db.commit()
        municipality_id = municipality.id
        selected_id = selected.id
        remaining_id = remaining.id

    csrf = login(client)
    page = client.get("/admin/personas")
    assert "data-confirm-selected" in page.text
    assert "data-select-all" in page.text
    deleted = client.post(
        "/admin/personas/eliminar",
        data={"csrf_token": csrf, "person_ids": [original_id, selected_id]},
        follow_redirects=False,
    )
    assert deleted.status_code == 303
    assert deleted.headers["location"] == "/admin/personas?eliminadas=2"
    with app.state.Session() as db:
        assert db.get(Person, original_id) is None
        assert db.get(Person, selected_id) is None
        assert db.get(Person, remaining_id) is not None
        assert db.get(Municipality, municipality_id) is not None
    assert client.get(f"/p/{token}").status_code == 404

    municipalities = client.get("/admin/municipios")
    assert f"/admin/municipios/{municipality_id}/personas/eliminar" in municipalities.text
    deleted_all = client.post(
        f"/admin/municipios/{municipality_id}/personas/eliminar",
        data={"csrf_token": csrf},
        follow_redirects=False,
    )
    assert deleted_all.status_code == 303
    assert deleted_all.headers["location"] == "/admin/municipios?eliminadas=1"
    with app.state.Session() as db:
        assert db.scalar(select(func.count()).select_from(Person).where(Person.municipality_id == municipality_id)) == 0
        assert db.get(Municipality, municipality_id) is not None


def test_delete_requires_csrf(test_app):
    _, client, _, person_id = test_app
    login(client)
    response = client.post(f"/admin/personas/{person_id}/eliminar", data={"csrf_token": "incorrecto"})
    assert response.status_code == 403
