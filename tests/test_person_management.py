from __future__ import annotations

import io
import json

import pytest
from openpyxl import load_workbook
from sqlalchemy import func, select

from app.gis import tile_xy, verify_generated_tiles
from app.models import Job, Municipality, Person, TileWindow
from tests.conftest import login


def test_admin_can_create_one_person_and_queue_tiles(test_app):
    app, client, _, _ = test_app
    csrf = login(client)
    with app.state.Session() as db:
        municipality_id = db.scalar(select(Municipality.id).where(Municipality.name == "San Rafael"))

    page = client.get("/admin/personas/nueva")
    assert page.status_code == 200
    assert "Crear una persona" in page.text
    response = client.post(
        "/admin/personas/nueva",
        data={
            "csrf_token": csrf,
            "municipality_id": municipality_id,
            "name": "  Beatriz   Gómez ",
            "phone": "+57 301 222 3344",
            "coordinates": "6.295,-75.032",
            "observation": "Visitar temprano",
        },
        follow_redirects=False,
    )
    assert response.status_code == 303
    with app.state.Session() as db:
        person = db.scalar(select(Person).where(Person.phone == "3012223344"))
        assert person is not None
        assert person.name == "Beatriz Gómez"
        assert person.reference_lon == -75.032
        assert person.reference_lat == 6.295
        job = db.scalar(select(Job).where(Job.job_type == "PREPARE_WINDOW", Job.payload_json == json.dumps({"person_id": person.id}, sort_keys=True, separators=(",", ":"))))
        assert job is not None
        assert job.status == "PENDIENTE"


def test_manual_person_rejects_duplicate_phone(test_app):
    app, client, _, _ = test_app
    csrf = login(client)
    with app.state.Session() as db:
        municipality_id = db.scalar(select(Municipality.id))
        before = db.scalar(select(func.count()).select_from(Person))
    response = client.post(
        "/admin/personas/nueva",
        data={"csrf_token": csrf, "municipality_id": municipality_id, "name": "Duplicada", "phone": "300 123 4567", "coordinates": "", "observation": ""},
    )
    assert response.status_code == 409
    assert "Ya existe una persona con ese celular" in response.text
    with app.state.Session() as db:
        assert db.scalar(select(func.count()).select_from(Person)) == before


def test_selected_and_municipality_tile_actions_queue_only_people_with_coordinates(test_app):
    app, client, _, person_id = test_app
    csrf = login(client)
    with app.state.Session() as db:
        municipality_id = db.scalar(select(Municipality.id))
        no_coordinates = Person(municipality_id=municipality_id, name="Sin coordenadas", phone="3008887777")
        db.add(no_coordinates)
        db.commit()
        no_coordinates_id = no_coordinates.id

    page = client.get("/admin/personas")
    assert "data-prepare-selected" in page.text
    queued = client.post(
        "/admin/personas/preparar-teselas",
        data={"csrf_token": csrf, "person_ids": [person_id, no_coordinates_id]},
        follow_redirects=False,
    )
    assert queued.status_code == 303
    assert "teselas_solicitadas=1" in queued.headers["location"]
    assert "sin_coordenadas=1" in queued.headers["location"]
    with app.state.Session() as db:
        assert db.scalar(select(func.count()).select_from(Job).where(Job.job_type == "PREPARE_WINDOW")) == 1

    queued_again = client.post(
        f"/admin/municipios/{municipality_id}/preparar-teselas",
        data={"csrf_token": csrf},
        follow_redirects=False,
    )
    assert queued_again.status_code == 303
    with app.state.Session() as db:
        assert db.scalar(select(func.count()).select_from(Job).where(Job.job_type == "PREPARE_WINDOW")) == 1


def test_downloaded_links_only_include_verified_images(test_app):
    app, client, _, person_id = test_app
    login(client)
    pending_sheet = load_workbook(io.BytesIO(client.get("/admin/enlaces").content), read_only=True).active
    assert list(pending_sheet.iter_rows(values_only=True)) == [("NOMBRE", "CELULAR", "MUNICIPIO", "ENLACE")]
    with app.state.Session() as db:
        person = db.get(Person, person_id)
        person.orthophoto_status = "LISTO"
        db.add(TileWindow(person_id=person_id, relative_dir="verified", min_zoom=15, max_zoom=19, status="LISTO"))
        for zoom in range(15, 20):
            x, y = tile_xy(person.reference_lon, person.reference_lat, zoom)
            tile = app.state.settings.data_dir / "tiles" / "verified" / str(zoom) / str(x) / f"{y}.png"
            tile.parent.mkdir(parents=True, exist_ok=True)
            tile.write_bytes(b"\x89PNG\r\n\x1a\nverified")
        db.commit()
    ready_sheet = load_workbook(io.BytesIO(client.get("/admin/enlaces").content), read_only=True).active
    assert len(list(ready_sheet.iter_rows(values_only=True))) == 2


def test_public_link_shows_preparation_message_until_tiles_are_ready(test_app):
    _, client, token, _ = test_app
    response = client.get(f"/p/{token}")
    assert response.status_code == 409
    assert "Estamos preparando la imagen" in response.text


def test_generated_tiles_are_verified_at_reference_point(tmp_path):
    lon, lat = -75.031, 6.294
    root = tmp_path / "tiles"
    for zoom in (15, 16):
        x, y = tile_xy(lon, lat, zoom)
        tile = root / str(zoom) / str(x) / f"{y}.png"
        tile.parent.mkdir(parents=True)
        tile.write_bytes(b"\x89PNG\r\n\x1a\nverified")
    assert verify_generated_tiles(root, lon, lat, 15, 16) == 2
    (root / "16").rename(root / "missing-16")
    with pytest.raises(RuntimeError, match="punto de referencia"):
        verify_generated_tiles(root, lon, lat, 15, 16)
