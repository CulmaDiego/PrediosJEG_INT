from __future__ import annotations

import io

from openpyxl import load_workbook

from app.config import Settings
from app.models import Person, TileWindow
from app.security import stable_public_token
from tests.conftest import login


def test_production_default_uses_external_hostname(monkeypatch):
    monkeypatch.setenv('PREDIOS_ENV', 'production')
    monkeypatch.delenv('PREDIOS_PUBLIC_BASE_URL', raising=False)
    settings = Settings.from_env()
    assert settings.public_base_url == 'https://interno.prediosjeg.site'


def test_external_hostname_is_used_in_individual_and_excel_links(test_app):
    app, client, _, person_id = test_app
    app.state.settings.public_base_url = 'https://interno.prediosjeg.site'
    expected = f'https://interno.prediosjeg.site/p/{stable_public_token(person_id, app.state.settings)}'
    base = app.state.settings.data_dir / 'tiles' / 'external-links'
    base.mkdir(parents=True)
    (base / 'source.vrt').write_text('<VRTDataset/>')
    with app.state.Session() as db:
        db.get(Person, person_id).orthophoto_status = 'LISTO'
        db.add(TileWindow(person_id=person_id, relative_dir='external-links', min_zoom=15, max_zoom=19,
                          render_mode='ON_DEMAND', status='LISTO'))
        db.commit()
    csrf = login(client)
    assert expected in client.get(f'/admin/personas/{person_id}').text
    for response in (client.get('/admin/enlaces'), client.post('/admin/municipios/1/enlaces', data={'csrf_token': csrf})):
        assert response.status_code == 200
        workbook = load_workbook(io.BytesIO(response.content), read_only=True)
        assert workbook.active.cell(2, 4).value == expected
        workbook.close()
