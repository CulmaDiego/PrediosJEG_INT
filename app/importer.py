from __future__ import annotations

import csv
import io
import json
import re
from pathlib import Path

from openpyxl import load_workbook


REQUIRED_COLUMNS = ("NOMBRE", "CELULAR", "COORDENADAS", "OBSERVACION")


def normalize_phone(value: object) -> str:
    digits = re.sub(r"\D", "", str(value or ""))
    if digits.startswith("57") and len(digits) == 12:
        digits = digits[2:]
    if not 7 <= len(digits) <= 15:
        raise ValueError("CELULAR debe contener entre 7 y 15 dígitos")
    return digits


def parse_coordinates(value: object) -> tuple[float, float] | tuple[None, None]:
    text = str(value or "").strip()
    if not text:
        return None, None
    try:
        if text.startswith("{"):
            point = json.loads(text)
            if point.get("type") != "Point" or len(point.get("coordinates", [])) != 2:
                raise ValueError
            lon, lat = map(float, point["coordinates"])
        else:
            parts = [part.strip() for part in re.split(r"[,;]", text)]
            if len(parts) != 2:
                raise ValueError
            lat, lon = map(float, parts)
        if not (-90 <= lat <= 90 and -180 <= lon <= 180):
            raise ValueError
        return lon, lat
    except (ValueError, TypeError, json.JSONDecodeError) as exc:
        raise ValueError("COORDENADAS debe ser 'latitud,longitud' o un Point GeoJSON") from exc


def _read_csv(content: bytes) -> list[dict[str, object]]:
    decoded = None
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            decoded = content.decode(encoding)
            break
        except UnicodeDecodeError:
            pass
    if decoded is None:
        raise ValueError("El CSV no está codificado en UTF-8 o Windows-1252")
    sample = decoded[:4096]
    dialect = csv.Sniffer().sniff(sample, delimiters=",;")
    return list(csv.DictReader(io.StringIO(decoded), dialect=dialect))


def _read_xlsx(path: Path) -> list[dict[str, object]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        iterator = sheet.iter_rows(values_only=True)
        headers = [str(value or "").strip().upper() for value in next(iterator, [])]
        return [dict(zip(headers, row, strict=False)) for row in iterator if any(value not in (None, "") for value in row)]
    finally:
        workbook.close()


def validate_import(path: Path, original_name: str) -> tuple[list[dict[str, object]], list[dict[str, object]]]:
    suffix = Path(original_name).suffix.lower()
    if suffix == ".csv":
        raw_rows = _read_csv(path.read_bytes())
    elif suffix == ".xlsx":
        raw_rows = _read_xlsx(path)
    else:
        raise ValueError("Solo se permiten archivos .csv y .xlsx")
    if not raw_rows:
        return [], [{"fila": 1, "error": "El archivo no contiene datos"}]
    columns = {str(column).strip().upper() for column in raw_rows[0]}
    missing = [column for column in REQUIRED_COLUMNS if column not in columns]
    if missing:
        return [], [{"fila": 1, "error": f"Faltan columnas: {', '.join(missing)}"}]

    valid: list[dict[str, object]] = []
    errors: list[dict[str, object]] = []
    phones: set[str] = set()
    for index, source in enumerate(raw_rows, start=2):
        row = {str(key).strip().upper(): value for key, value in source.items()}
        row_errors: list[str] = []
        name = str(row.get("NOMBRE") or "").strip()
        if not name or len(name) > 200:
            row_errors.append("NOMBRE es obligatorio y admite máximo 200 caracteres")
        try:
            phone = normalize_phone(row.get("CELULAR"))
            if phone in phones:
                row_errors.append("CELULAR está repetido en el archivo")
            phones.add(phone)
        except ValueError as exc:
            phone = ""
            row_errors.append(str(exc))
        try:
            lon, lat = parse_coordinates(row.get("COORDENADAS"))
        except ValueError as exc:
            lon, lat = None, None
            row_errors.append(str(exc))
        observation = str(row.get("OBSERVACION") or "").strip()
        if len(observation) > 2000:
            row_errors.append("OBSERVACION admite máximo 2000 caracteres")
        if row_errors:
            errors.append({"fila": index, "error": "; ".join(row_errors)})
        else:
            valid.append(
                {
                    "row": index,
                    "name": name,
                    "phone": phone,
                    "reference_lon": lon,
                    "reference_lat": lat,
                    "observation": observation,
                    "metadata": {},
                }
            )
    return valid, errors

